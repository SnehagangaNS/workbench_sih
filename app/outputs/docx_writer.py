"""
docx_writer.py
--------------
Generates executive-ready .docx deliverables with professional styling,
metadata banners, section callouts, and RAG source citations.
"""

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT_DIR = Path(__file__).parent.parent.parent / "data" / "outputs"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def set_cell_shading(cell, color_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def build_report(
    title: str,
    sections: list[dict],  # [{"heading": str, "body": str}, ...]
    sources: list[str] | None = None,
    filename: str = "report.docx",
) -> str:
    doc = Document()

    # Set 1-inch margins
    sections_env = doc.sections
    for sec in sections_env:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run = title_p.add_run(title)
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)  # Slate dark

    # Subtitle / Accent Line
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    sub_run = sub_p.add_run("Executive Document Summary & Analysis")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(12)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)  # Sky blue accent

    # Metadata Block Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_00 = meta_table.cell(0, 0)
    cell_01 = meta_table.cell(0, 1)
    cell_10 = meta_table.cell(1, 0)
    cell_11 = meta_table.cell(1, 1)

    set_cell_shading(cell_00, "F8FAFC")
    set_cell_shading(cell_01, "F8FAFC")
    set_cell_shading(cell_10, "F8FAFC")
    set_cell_shading(cell_11, "F8FAFC")

    cell_00.paragraphs[0].add_run("Generated: ").bold = True
    cell_00.paragraphs[0].add_run(datetime.now().strftime("%B %d, %Y - %H:%M"))
    cell_01.paragraphs[0].add_run("Status: ").bold = True
    cell_01.paragraphs[0].add_run("Final Deliverable")
    cell_10.paragraphs[0].add_run("Platform: ").bold = True
    cell_10.paragraphs[0].add_run("Local Offline AI Workbench")
    cell_11.paragraphs[0].add_run("Classification: ").bold = True
    cell_11.paragraphs[0].add_run("Internal Analysis")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Render Sections
    for idx, section in enumerate(sections):
        heading_text = section.get("heading", f"Section {idx + 1}")
        body_text = section.get("body", "")

        # Section Heading
        h_p = doc.add_paragraph()
        h_p.paragraph_format.space_before = Pt(14)
        h_p.paragraph_format.space_after = Pt(6)
        h_p.paragraph_format.keep_with_next = True
        h_run = h_p.add_run(heading_text)
        h_run.font.name = "Calibri"
        h_run.font.size = Pt(16)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        # Section Body Paragraphs
        for para in body_text.split("\n\n"):
            p_clean = para.strip()
            if not p_clean:
                continue

            # Check if paragraph is bullet list item
            if p_clean.startswith("- ") or p_clean.startswith("* "):
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run(p_clean[2:])
                run.font.name = "Calibri"
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.line_spacing = 1.15
                run = p.add_run(p_clean)
                run.font.name = "Calibri"
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # RAG Sources Citation Section
    if sources:
        h_p = doc.add_paragraph()
        h_p.paragraph_format.space_before = Pt(18)
        h_p.paragraph_format.space_after = Pt(6)
        h_run = h_p.add_run("Referenced Sources & Citations")
        h_run.font.name = "Calibri"
        h_run.font.size = Pt(14)
        h_run.font.bold = True
        h_run.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

        for src in sources:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(src)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.italic = True

    out_path = OUTPUT_DIR / filename
    doc.save(str(out_path))
    return str(out_path)
