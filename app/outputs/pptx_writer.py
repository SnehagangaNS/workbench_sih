"""
pptx_writer.py
--------------
Generates clean, structured PowerPoint (.pptx) presentation decks.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

OUTPUT_DIR = Path(__file__).parent.parent.parent / "data" / "outputs"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def build_deck(
    title: str,
    subtitle: str,
    slides: list[dict],  # [{"heading": str, "bullets": [str, ...]}, ...]
    filename: str = "deck.pptx",
) -> str:
    prs = Presentation()

    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title_shape = slide.shapes.title
    subtitle_shape = slide.placeholders[1]

    title_shape.text = title
    title_tf = title_shape.text_frame
    if title_tf.paragraphs:
        p = title_tf.paragraphs[0]
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    subtitle_shape.text = subtitle
    sub_tf = subtitle_shape.text_frame
    if sub_tf.paragraphs:
        p = sub_tf.paragraphs[0]
        p.font.size = Pt(20)
        p.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

    # Content slides
    bullet_layout = prs.slide_layouts[1]
    for s in slides:
        slide = prs.slides.add_slide(bullet_layout)
        heading_shape = slide.shapes.title
        heading_shape.text = s.get("heading", "Key Insights")
        if heading_shape.text_frame.paragraphs:
            p = heading_shape.text_frame.paragraphs[0]
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()
        bullets = s.get("bullets", [])
        for i, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = bullet
            p.font.size = Pt(16)
            p.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            p.space_after = Pt(10)

    out_path = OUTPUT_DIR / filename
    prs.save(str(out_path))
    return str(out_path)
